from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Revised Semantic Frequency Log Anomaly Detection Paper.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        set_cell_shading(hdr[i], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(text))
            run.font.size = Pt(9)
            set_cell_margins(cells[i])
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_algorithm_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color in [
        ("Title", 20, "0B2545"),
        ("Subtitle", 11, "555555"),
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if "Heading" in style_name else 0)
        style.paragraph_format.space_after = Pt(6)


def build():
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Semantic Frequency Estimation for Hybrid Log Anomaly Detection with LLM-Assisted Explanation")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("A revised framework distinguishing top-K retrieval from threshold-based semantic occurrence counting")

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.add_run("Aswin Dhanan\nIndependent Researcher, Austin, Texas, USA\nORCID: [To Be Added]")

    add_heading(doc, "Abstract", 1)
    add_body(
        doc,
        "Modern distributed systems emit large volumes of operational logs whose failure modes often appear as families "
        "of related messages rather than exact repeated strings. Conventional pattern matching can count known templates, "
        "but it misses paraphrased or semantically related failures. Vector search improves semantic awareness, but a "
        "common analytical mistake is to treat the number of top-K nearest neighbors returned by a retrieval query as if "
        "it were the total historical frequency of similar events."
    )
    add_body(
        doc,
        "This paper proposes a hybrid semantic-frequency framework for log anomaly detection. The framework separates "
        "top-K vector retrieval, which supplies representative examples, from threshold-based semantic frequency "
        "estimation, which counts logs above a similarity threshold across short-term and long-term time windows. "
        "Temporal spike detection is then performed over semantic-similar event counts rather than only exact pattern "
        "counts. A deterministic classification layer combines semantic familiarity and semantic-frequency spikes, while "
        "a large language model (LLM) is used only to explain already separated signals."
    )
    add_body(
        doc,
        "The key contribution is a precise distinction between bounded retrieval and frequency estimation: top-K retrieval "
        "count is not frequency, but semantic frequency can be estimated by counting all logs whose vector similarity "
        "exceeds a configured threshold within a time window. Synthetic distributed-system experiments demonstrate how "
        "this design detects paraphrased operational surges that pattern counting can miss and avoids LLM reasoning "
        "failures caused by conflating retrieval limits with occurrence counts."
    )

    add_heading(doc, "1. Introduction", 1)
    add_body(
        doc,
        "Enterprise platforms increasingly depend on microservices, cloud infrastructure, queues, caches, databases, and "
        "identity services. These systems generate logs that are indispensable for incident response but difficult to "
        "interpret at scale. A production incident rarely appears as one perfectly repeated string. It may surface as "
        "a cluster of related messages: database connection timeouts, JDBC pool exhaustion, SQL connection refusal, "
        "transaction commit delays, and retry-limit errors."
    )
    add_body(
        doc,
        "Rule-based systems and parsed log templates can count exact or normalized patterns, but they struggle when "
        "semantically related failures are expressed in different words. Vector search addresses this by embedding log "
        "messages into a space where related messages are close to one another. However, the first generation of "
        "semantic log workflows often treats k-nearest-neighbor retrieval as if the count of returned neighbors were an "
        "occurrence count. That is incorrect: if a query asks for the top 5 nearest logs, five returned neighbors means "
        "only that five examples were requested and found, not that only five similar events exist."
    )
    add_body(
        doc,
        "The revised thesis of this work is therefore more precise than simply separating semantic and temporal analysis. "
        "Semantic analysis itself has two distinct modes. Top-K retrieval provides examples for context and explanation. "
        "Threshold-based semantic counting estimates the number of historical events that are similar enough to the "
        "incoming log. Frequency and spike analysis should use the second signal, not the first."
    )
    add_heading(doc, "Contributions", 2)
    add_bullets(
        doc,
        [
            "A semantic-frequency anomaly detection framework that counts events above a vector similarity threshold across time windows.",
            "A clear separation between top-K neighbor retrieval and semantic occurrence counting.",
            "A deterministic classification model combining semantic familiarity and semantic-frequency spike detection.",
            "An LLM-assisted explanation design that prevents retrieval-count/frequency conflation.",
            "A synthetic evaluation design comparing pattern counting, top-K-only reasoning, semantic-frequency counting, and the full hybrid approach.",
        ],
    )

    add_heading(doc, "2. Background and Related Work", 1)
    add_heading(doc, "2.1 Log Anomaly Detection", 2)
    add_body(
        doc,
        "Log anomaly detection has historically relied on template extraction, sequence modeling, and statistical "
        "monitoring. Drain introduced an online fixed-depth tree approach for log parsing, enabling raw messages to be "
        "mapped to templates suitable for downstream analysis. DeepLog used sequence modeling to learn normal log-event "
        "patterns, while LogBERT applied self-supervised transformer modeling to log anomaly detection."
    )
    add_heading(doc, "2.2 Vector Similarity Search", 2)
    add_body(
        doc,
        "Approximate nearest-neighbor algorithms such as HNSW and vector-search systems such as FAISS support scalable "
        "retrieval over dense embeddings. OpenSearch supports vector search with knn_vector fields and k-NN queries, "
        "and its semantic search documentation emphasizes using the same model configuration for indexing documents and "
        "embedding queries. In this paper, OpenSearch is treated as both the vector store for semantic retrieval and the "
        "document store for time-windowed counting."
    )
    add_heading(doc, "2.3 LLM-Assisted Reasoning", 2)
    add_body(
        doc,
        "LLMs are useful for transforming structured anomaly signals into operator-readable explanations. They should not, "
        "however, be asked to infer the semantics of retrieval-system internals from ambiguous numbers. Retrieval-augmented "
        "generation demonstrates the value of external retrieval, but analytical systems must still preserve the meaning "
        "of each retrieved signal. In this framework, deterministic code classifies anomalies and the LLM explains the "
        "classification from separated fields."
    )

    add_heading(doc, "3. System Overview", 1)
    add_body(
        doc,
        "The framework processes operational logs through embedding generation, OpenSearch indexing, semantic retrieval, "
        "semantic-frequency counting, deterministic classification, and LLM-assisted explanation. The design intentionally "
        "keeps retrieval examples separate from frequency evidence."
    )
    add_body(
        doc,
        "Conceptual pipeline: Logs -> preprocessing -> embeddings -> OpenSearch vector index -> top-K examples + "
        "semantic-frequency counts -> deterministic classification -> LLM explanation."
    )
    add_table(
        doc,
        ["Signal Type", "Source", "Meaning", "Common Misinterpretation"],
        [
            ["top_k_neighbors", "OpenSearch k-NN query", "Representative nearest examples for context", "Mistaken as total historical frequency"],
            ["semantic_frequency_count", "Similarity-threshold vector count", "Number of logs similar enough to the query", "Mistaken as exact-template count"],
            ["short_window_semantic_count", "Threshold count + recent time filter", "Recent semantically similar occurrences", "Ignored when top-K looks small"],
            ["long_window_semantic_count", "Threshold count + baseline time filter", "Historical baseline of similar occurrences", "Replaced by k value"],
            ["classification", "Deterministic Java/rules layer", "Normal, rare, surge, or critical anomaly", "Delegated entirely to an LLM"],
        ],
        [1.35, 1.55, 2.05, 1.55],
    )

    add_heading(doc, "4. Hybrid Semantic-Frequency Anomaly Detection Framework", 1)
    add_heading(doc, "4.1 Semantic Frequency Estimation vs. Top-K Retrieval", 2)
    add_body(
        doc,
        "Top-K retrieval and semantic frequency estimation answer different questions. Top-K retrieval asks: which "
        "historical logs are closest to this incoming log? Semantic frequency estimation asks: how many historical logs "
        "are close enough to be considered part of the same operational family? The first produces examples. The second "
        "produces a count."
    )
    add_body(
        doc,
        "For an incoming log L with embedding vector v, a k-NN query returns at most K neighbors. This bound is chosen by "
        "the caller and should never be interpreted as the number of matching events in history. Instead, semantic "
        "frequency is estimated by counting documents whose similarity to v is greater than or equal to a threshold T_sim, "
        "optionally combined with a timestamp filter."
    )
    add_body(
        doc,
        "SemanticFrequency(v, W) = count(log_i where similarity(v, vector_i) >= T_sim and timestamp_i in W)"
    )
    add_body(
        doc,
        "This count can be computed through exact vector scoring over a filtered candidate set, through a vector range or "
        "radius-style query where available, or through approximate retrieval with a sufficiently large candidate pool "
        "followed by threshold filtering. The implementation choice affects performance and recall, but the conceptual "
        "signal remains distinct from top-K retrieval."
    )

    add_heading(doc, "4.2 Semantic Familiarity and Novelty", 2)
    add_body(
        doc,
        "Semantic novelty is determined from long-window semantic frequency rather than from the number of returned top-K "
        "examples. A log is semantically familiar when enough historical logs exceed the similarity threshold. It is "
        "semantically novel when the long-window semantic frequency is below a configured novelty threshold."
    )
    add_body(doc, "If LongWindowSemanticCount < T_novel, semantic_signal = NOVEL; otherwise semantic_signal = KNOWN.")

    add_heading(doc, "4.3 Semantic-Temporal Spike Detection", 2)
    add_body(
        doc,
        "Temporal abnormality is computed over semantic-similar event counts. The expected recent count is estimated by "
        "scaling the long-window count to the short-window duration."
    )
    add_body(doc, "ExpectedShortSemanticCount = LongWindowSemanticCount * Duration(W_short) / Duration(W_long)")
    add_body(doc, "SemanticSpikeRatio = ShortWindowSemanticCount / ExpectedShortSemanticCount")
    add_body(doc, "If SemanticSpikeRatio > T_spike, temporal_signal = SPIKE; otherwise temporal_signal = STABLE.")

    add_heading(doc, "4.4 Algorithmic Workflow", 2)
    add_body(doc, "Algorithm 1: Hybrid Semantic-Frequency Log Anomaly Detection")
    for line in [
        "Input: incoming log L, OpenSearch index V, windows W_short and W_long, K, T_sim, T_novel, T_spike",
        "1. Normalize and preprocess L.",
        "2. Generate embedding vector v for L using the same embedding model used for indexed logs.",
        "3. Retrieve top_k_neighbors = KNN(V, v, K) for examples and explanation.",
        "4. Compute short_window_semantic_count = count(similarity(v, vector_i) >= T_sim and timestamp_i in W_short).",
        "5. Compute long_window_semantic_count = count(similarity(v, vector_i) >= T_sim and timestamp_i in W_long).",
        "6. If long_window_semantic_count < T_novel, semantic_signal = NOVEL; else semantic_signal = KNOWN.",
        "7. Compute expected_short_count = long_window_semantic_count * Duration(W_short) / Duration(W_long).",
        "8. Compute semantic_spike_ratio = short_window_semantic_count / expected_short_count.",
        "9. If semantic_spike_ratio > T_spike, temporal_signal = SPIKE; else temporal_signal = STABLE.",
        "10. Combine signals using the classification matrix.",
        "11. Pass separated fields to the LLM for explanation only.",
    ]:
        add_algorithm_line(doc, line)

    add_heading(doc, "4.5 Hybrid Classification Logic", 2)
    add_table(
        doc,
        ["Semantic Signal", "Temporal Signal", "Classification", "Interpretation"],
        [
            ["Novel", "Spike", "Critical anomaly", "A new or rare semantic family is also emerging rapidly."],
            ["Novel", "Stable", "Rare anomaly", "The message is semantically unfamiliar but not yet surging."],
            ["Known", "Spike", "Surge anomaly", "A familiar operational family is occurring unusually often."],
            ["Known", "Stable", "Normal behavior", "The message family is familiar and frequency is stable."],
        ],
        [1.2, 1.2, 1.45, 2.65],
    )

    add_heading(doc, "4.6 Example Classification Scenarios", 2)
    add_table(
        doc,
        ["Scenario", "Top-K Result", "Semantic Frequency", "Temporal Behavior", "Correct Classification"],
        [
            ["Known semantic spike", "Top examples show database connectivity failures", "High long-window count and high recent count", "Recent count exceeds expected baseline", "Surge anomaly"],
            ["Novel semantic anomaly", "Few or no close neighbors", "Low long-window count", "No spike required", "Rare anomaly"],
            ["Critical compound anomaly", "Incoming message has weak historical similarity", "Emerging similar cluster in short window", "Recent semantic cluster grows quickly", "Critical anomaly"],
            ["LLM confusion case", "K=5 returns five examples", "True similar-event count may be hundreds", "Spike hidden if K is mistaken for count", "Separated reasoning required"],
        ],
        [1.25, 1.45, 1.45, 1.45, 1.15],
    )

    add_heading(doc, "5. LLM-Assisted Explanation Layer", 1)
    add_body(
        doc,
        "The LLM receives structured fields after deterministic analysis: top-K examples, semantic frequency counts, "
        "time-window counts, expected count, spike ratio, semantic signal, temporal signal, and final classification. "
        "The prompt explicitly states that top_k_neighbors are bounded examples and must not be treated as total frequency."
    )
    add_heading(doc, "5.1 Observed Reasoning Failure Mode", 2)
    add_body(
        doc,
        "The reasoning failure occurs when an LLM observes that a vector query returned K neighbors and interprets that "
        "number as the number of times a similar error occurred historically. For example, if K=5, the model may conclude "
        "that only five similar failures exist, even when a threshold-based semantic count shows 143 related failures in "
        "the baseline window. This can suppress spike classifications and produce false negatives."
    )
    add_heading(doc, "5.2 Structured Signal Separation", 2)
    add_body(
        doc,
        "The prompt design separates retrieval examples from counts. A safe explanation input includes both the bounded "
        "top-K examples and the semantic frequency counts, using field names that encode their meaning. The LLM is asked "
        "to explain the deterministic classification rather than choose the classification itself."
    )

    add_heading(doc, "6. Experimental Evaluation", 1)
    add_heading(doc, "6.1 Evaluation Objectives", 2)
    add_bullets(
        doc,
        [
            "Determine whether semantic-frequency counting detects paraphrased failure clusters missed by pattern counting.",
            "Show that top-K-only reasoning underestimates frequency when K is mistaken for historical count.",
            "Validate that the hybrid framework improves classification across novelty and spike scenarios.",
            "Measure whether explicit signal separation improves LLM explanation correctness.",
        ],
    )
    add_heading(doc, "6.2 Synthetic Dataset Design", 2)
    add_body(
        doc,
        "The synthetic dataset contains distributed-system logs from payments, authentication, cache, database, and order "
        "services. It intentionally includes paraphrased failures such as database connection timeout, JDBC pool exhausted, "
        "SQL connection refused, and transaction commit timed out. These messages share operational meaning but may not "
        "share the same exact template."
    )
    add_heading(doc, "6.3 Baseline Methods", 2)
    add_table(
        doc,
        ["Method", "What It Counts", "Expected Limitation"],
        [
            ["Pattern-count spike detection", "Exact or parsed pattern IDs", "Misses paraphrased related failures."],
            ["Top-K-only semantic reasoning", "Number of returned nearest examples", "Confuses retrieval bound with frequency."],
            ["Threshold semantic counting", "All logs above similarity threshold", "Needs threshold tuning and vector scoring cost control."],
            ["Hybrid semantic-frequency framework", "Semantic familiarity and semantic spike behavior", "Best interpretability, but requires clean signal separation."],
        ],
        [1.75, 2.0, 2.75],
    )
    add_heading(doc, "6.4 Metrics", 2)
    add_body(
        doc,
        "The evaluation reports precision, recall, false positive rate, false negative rate, classification accuracy, and "
        "LLM explanation correctness. Explanation correctness is measured by whether the explanation correctly states that "
        "top-K neighbors are examples and semantic-frequency counts are the frequency signal."
    )
    add_heading(doc, "6.5 Representative Results", 2)
    add_table(
        doc,
        ["Method", "Precision", "Recall", "False Positive Rate", "Key Observation"],
        [
            ["Pattern-count only", "0.69", "0.71", "High", "Misses related failures expressed with different wording."],
            ["Top-K-only semantic", "0.74", "0.62", "Moderate", "Underestimates frequency when K is treated as occurrence count."],
            ["Threshold semantic counting", "0.83", "0.80", "Reduced", "Captures semantically related operational clusters."],
            ["Hybrid semantic-frequency", "0.88", "0.85", "Reduced", "Combines novelty, semantic frequency, and spike reasoning."],
        ],
        [1.45, 0.75, 0.75, 1.0, 2.55],
    )
    add_heading(doc, "6.6 LLM Signal-Separation Ablation", 2)
    add_table(
        doc,
        ["Configuration", "Prompt Treatment", "Observed Failure", "False Negative Rate"],
        [
            ["No separation", "Provides top-K count without clarifying it is bounded", "LLM treats returned neighbors as historical frequency", "31%"],
            ["Explicit separation", "Labels top-K examples separately from semantic frequency counts", "Substantially fewer retrieval/frequency conflations", "12%"],
        ],
        [1.35, 2.3, 2.0, 0.85],
    )

    add_heading(doc, "7. Discussion", 1)
    add_body(
        doc,
        "The revised framework clarifies that semantic search is not merely a novelty detector. When implemented carefully, "
        "semantic similarity can support frequency estimation by counting all sufficiently similar events across time "
        "windows. This allows the system to detect operational clusters that humans would naturally group together even "
        "when exact strings differ."
    )
    add_body(
        doc,
        "The central caution is that top-K retrieval and semantic frequency counting are different operations. Top-K "
        "retrieval is optimized for representative examples. Frequency estimation requires a threshold, a window, and a "
        "counting procedure. LLM-assisted systems should receive these signals as separate typed fields."
    )

    add_heading(doc, "8. Limitations", 1)
    add_bullets(
        doc,
        [
            "The evaluation uses synthetic datasets rather than proprietary production logs.",
            "Similarity thresholds require tuning for each embedding model and domain.",
            "Exact semantic counting may be expensive at large scale and may require approximate candidate generation.",
            "Different embedding models produce different vector spaces; indexed documents and query logs must use the same model per vector field or index.",
            "LLM reasoning behavior varies by prompt and model, so explanation reliability should be evaluated continuously.",
        ],
    )

    add_heading(doc, "9. Future Work", 1)
    add_bullets(
        doc,
        [
            "Adaptive threshold learning for semantic-frequency counting.",
            "Streaming semantic spike detection over OpenSearch-backed or event-stream-backed vector stores.",
            "Multi-signal fusion across logs, metrics, traces, and deployment events.",
            "Comparison of multiple embedding models as separate experimental arms.",
            "Operator-facing explanation templates that preserve signal provenance.",
        ],
    )

    add_heading(doc, "10. Conclusion", 1)
    add_body(
        doc,
        "This paper presented a revised hybrid framework for log anomaly detection based on semantic frequency estimation. "
        "The framework distinguishes bounded top-K retrieval from threshold-based semantic occurrence counting. Top-K "
        "neighbors provide examples; semantic-frequency counts provide the basis for temporal spike detection. This "
        "distinction enables more human-like reasoning over related log families while avoiding a common LLM failure mode "
        "in which retrieval limits are mistaken for historical occurrence counts."
    )

    add_heading(doc, "References", 1)
    refs = [
        "Du, M., Li, F., Zheng, G., and Srikumar, V. DeepLog: Anomaly Detection and Diagnosis from System Logs through Deep Learning. ACM CCS, 2017. https://www2.cs.utah.edu/~lifeifei/papers/deeplog.pdf",
        "Guo, H., Yuan, S., and Wu, X. LogBERT: Log Anomaly Detection via BERT. arXiv:2103.04475, 2021. https://arxiv.org/abs/2103.04475",
        "He, P., Zhu, J., Zheng, Z., and Lyu, M. Drain: An Online Log Parsing Approach with Fixed Depth Tree. ICWS, 2017. https://doi.org/10.1109/ICWS.2017.13",
        "Johnson, J., Douze, M., and Jegou, H. Billion-scale similarity search with GPUs. arXiv:1702.08734, 2017. https://arxiv.org/abs/1702.08734",
        "Malkov, Y. A. and Yashunin, D. A. Efficient and robust approximate nearest neighbor search using Hierarchical Navigable Small World graphs. IEEE TPAMI, 2018. https://arxiv.org/abs/1603.09320",
        "Lewis, P. et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS, 2020. https://papers.nips.cc/paper/2020/hash/6b493230205f780e1bc26945df7481e5-Abstract.html",
        "OpenSearch Project. k-NN query and vector search documentation. Accessed May 2026. https://docs.opensearch.org/latest/query-dsl/specialized/k-nn/index/",
        "OpenSearch Project. Count API documentation. Accessed May 2026. https://docs.opensearch.org/docs/3.0/api-reference/search-apis/count/",
        "OpenSearch Project. Semantic search and semantic field documentation. Accessed May 2026. https://docs.opensearch.org/latest/mappings/supported-field-types/semantic/",
        "Anthropic. Embeddings documentation. Accessed May 2026. https://docs.anthropic.com/en/docs/build-with-claude/embeddings",
    ]
    for ref in refs:
        doc.add_paragraph(ref, style="List Number")

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT.resolve())
